import pandas as pd
import numpy as np
import io
import base64
from datetime import datetime
import xlsxwriter
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference, LineChart, ScatterChart
from openpyxl.utils.dataframe import dataframe_to_rows
from typing import Optional

class ExportManager:
    """Handles data export functionality including CSV and Excel with charts."""
    
    def to_csv(self, df: pd.DataFrame) -> str:
        """Convert DataFrame to CSV string."""
        return df.to_csv(index=False)
    
    def to_excel_with_charts(self, df: pd.DataFrame, dataset_name: str) -> bytes:
        """Create Excel file with data and charts."""
        buffer = io.BytesIO()
        
        with pd.ExcelWriter(buffer, engine='xlsxwriter') as writer:
            # Write main data
            df.to_excel(writer, sheet_name='Data', index=False)
            
            # Write summary statistics
            numeric_df = df.select_dtypes(include=[np.number])
            if len(numeric_df.columns) > 0:
                summary_stats = numeric_df.describe()
                summary_stats.to_excel(writer, sheet_name='Summary Statistics')
            
            # Write missing values analysis
            missing_counts = df.isnull().sum()
            missing_df = pd.DataFrame({
                'Column': missing_counts.index,
                'Missing Count': missing_counts.values,
                'Missing Percentage': (missing_counts / len(df) * 100).values
            })
            missing_df.to_excel(writer, sheet_name='Missing Values', index=False)
            
            # Get workbook and worksheet objects
            workbook = writer.book
            
            # Add charts if we have numeric data
            if len(numeric_df.columns) > 0:
                self._add_charts_to_workbook(workbook, df, numeric_df)
        
        buffer.seek(0)
        return buffer.read()
    
    def _add_charts_to_workbook(self, workbook, df: pd.DataFrame, numeric_df: pd.DataFrame):
        """Add charts to Excel workbook."""
        # Create charts worksheet
        charts_worksheet = workbook.add_worksheet('Charts')
        
        # Chart 1: Bar chart of column means
        if len(numeric_df.columns) > 0:
            means = numeric_df.mean()
            
            # Write data for chart
            charts_worksheet.write('A1', 'Column Means')
            charts_worksheet.write('A2', 'Column')
            charts_worksheet.write('B2', 'Mean Value')
            
            for idx, (col, mean_val) in enumerate(means.items()):
                charts_worksheet.write(idx + 3, 0, col)
                charts_worksheet.write(idx + 3, 1, mean_val)
            
            # Create bar chart
            chart = workbook.add_chart({'type': 'bar'})
            chart.add_series({
                'name': 'Mean Values',
                'categories': f'Charts!A4:A{3 + len(means)}',
                'values': f'Charts!B4:B{3 + len(means)}',
            })
            chart.set_title({'name': 'Mean Values by Column'})
            chart.set_x_axis({'name': 'Mean Value'})
            chart.set_y_axis({'name': 'Columns'})
            
            charts_worksheet.insert_chart('D2', chart)
    
    def create_summary_report(self, df: pd.DataFrame, dataset_name: str) -> str:
        """Create a text summary report."""
        report = []
        report.append(f"Data Analysis Report: {dataset_name}")
        report.append("=" * 50)
        report.append(f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        report.append("")
        
        # Basic info
        report.append("DATASET OVERVIEW")
        report.append("-" * 20)
        report.append(f"Rows: {df.shape[0]:,}")
        report.append(f"Columns: {df.shape[1]:,}")
        report.append(f"Memory Usage: {df.memory_usage().sum() / 1024**2:.2f} MB")
        report.append("")
        
        # Column types
        report.append("COLUMN TYPES")
        report.append("-" * 20)
        type_counts = df.dtypes.value_counts()
        for dtype, count in type_counts.items():
            report.append(f"{dtype}: {count} columns")
        report.append("")
        
        # Missing values
        missing_counts = df.isnull().sum()
        missing_total = missing_counts.sum()
        if missing_total > 0:
            report.append("MISSING VALUES")
            report.append("-" * 20)
            report.append(f"Total missing values: {missing_total:,}")
            report.append(f"Percentage of dataset: {(missing_total / (df.shape[0] * df.shape[1]) * 100):.2f}%")
            report.append("")
            report.append("Missing values by column:")
            for col, count in missing_counts[missing_counts > 0].items():
                percentage = (count / df.shape[0]) * 100
                report.append(f"  {col}: {count:,} ({percentage:.2f}%)")
            report.append("")
        
        # Numeric summary
        numeric_df = df.select_dtypes(include=[np.number])
        if len(numeric_df.columns) > 0:
            report.append("NUMERIC COLUMNS SUMMARY")
            report.append("-" * 20)
            summary = numeric_df.describe()
            for col in numeric_df.columns:
                report.append(f"{col}:")
                report.append(f"  Mean: {summary.loc['mean', col]:.3f}")
                report.append(f"  Median: {summary.loc['50%', col]:.3f}")
                report.append(f"  Std Dev: {summary.loc['std', col]:.3f}")
                report.append(f"  Min: {summary.loc['min', col]:.3f}")
                report.append(f"  Max: {summary.loc['max', col]:.3f}")
                report.append("")
        
        # Categorical summary
        categorical_df = df.select_dtypes(include=['object', 'category'])
        if len(categorical_df.columns) > 0:
            report.append("CATEGORICAL COLUMNS SUMMARY")
            report.append("-" * 20)
            for col in categorical_df.columns:
                unique_count = df[col].nunique()
                most_frequent = df[col].mode()
                most_frequent_value = most_frequent.iloc[0] if len(most_frequent) > 0 else "N/A"
                most_frequent_count = df[col].value_counts().iloc[0] if len(df[col].value_counts()) > 0 else 0
                
                report.append(f"{col}:")
                report.append(f"  Unique values: {unique_count}")
                report.append(f"  Most frequent: '{most_frequent_value}' ({most_frequent_count} times)")
                report.append("")
        
        return "\n".join(report)
    
    def to_json(self, df: pd.DataFrame) -> str:
        """Convert DataFrame to JSON string."""
        return df.to_json(orient='records', indent=2)
    
    def create_filtered_report(self, original_df: pd.DataFrame, filtered_df: pd.DataFrame) -> str:
        """Create a report comparing original and filtered datasets."""
        report = []
        report.append("FILTERING REPORT")
        report.append("=" * 20)
        report.append(f"Original dataset: {original_df.shape[0]:,} rows, {original_df.shape[1]} columns")
        report.append(f"Filtered dataset: {filtered_df.shape[0]:,} rows, {filtered_df.shape[1]} columns")
        report.append(f"Rows removed: {original_df.shape[0] - filtered_df.shape[0]:,}")
        report.append(f"Percentage retained: {(filtered_df.shape[0] / original_df.shape[0] * 100):.2f}%")
        report.append("")
        
        # Compare numeric summaries
        numeric_cols = original_df.select_dtypes(include=[np.number]).columns
        if len(numeric_cols) > 0:
            report.append("NUMERIC COLUMNS COMPARISON")
            report.append("-" * 30)
            
            orig_summary = original_df[numeric_cols].describe()
            filt_summary = filtered_df[numeric_cols].describe() if len(filtered_df) > 0 else pd.DataFrame()
            
            for col in numeric_cols:
                if col in filt_summary.columns:
                    report.append(f"{col}:")
                    report.append(f"  Original mean: {orig_summary.loc['mean', col]:.3f}")
                    report.append(f"  Filtered mean: {filt_summary.loc['mean', col]:.3f}")
                    mean_change = ((filt_summary.loc['mean', col] - orig_summary.loc['mean', col]) / 
                                 orig_summary.loc['mean', col] * 100)
                    report.append(f"  Mean change: {mean_change:.2f}%")
                    report.append("")
        
        return "\n".join(report)
